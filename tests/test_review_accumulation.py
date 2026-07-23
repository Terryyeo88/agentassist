"""tests/test_review_accumulation.py -- t-review-accumulation (D-2026-07-23-review-accumulation)
FAILING-FIRST.

BUILD ID: t-review-accumulation. Written BEFORE the API wiring exists; these MUST fail today
for the RIGHT reason. Phase 2 will:

  1. WIRE the ALREADY-PRESENT store (``agent/review_store.py`` exists in this worktree -- see the
     GAP note below) into the API:
       * NEW ``POST /review-session`` (optional {"label": str}) -> 200 {"review_id", "created_at"};
       * NEW ``GET /review-session/{id}`` -> 200 grouped-slices merged view (R6 -- NO flattened
         queue) {"review_id","created_at","label","slices":[...],"superseded":[...],
         "coverage_matrix":{check: [{"source_kind","level","reason"}, ...]},
         "validation_status":"unvalidated","disclaimer": str}; 404 unknown id;
       * NEW ``GET /review-session`` (list) -> {"reviews": [{"review_id","created_at","label",
         "slice_count","source_kinds"}]}.
  2. ADD an OPTIONAL multipart field ``review_id: str = Form(None)`` to ``post_review_upload``:
       * ABSENT -> byte-identical stateless behaviour (R1: response equal to a control upload
         without the field; NO new response keys EVER -- the 5/6/7-key literals stay);
       * present + unknown id -> 404; present + invalid chars ("../evil"/"UPPER") -> 422;
       * present + valid -> the branch response is UNCHANGED and a slice record is appended:
         {source_kind, sha256 (of the uploaded primary file bytes), uploaded_at, client_id
         (record-only: xero_f5_upload->xero_demo, xero_sales_upload->xero_sales_demo,
         extract_review->extract_demo), period (may be null), queue rows, coverage_status rows,
         + branch extras}.
  3. RE-UPLOAD SEMANTICS (R2): same sha256 + same source_kind -> IDEMPOTENT SKIP (no new slice);
     different sha256 + same source_kind -> SUPERSEDE-IN-VIEW (new slice active, old moves to
     "superseded" with superseded_at; the store JSONL keeps BOTH lines -- append-only).

THREE-TIMES RULE (prompt + code + THIS test): the "uploads ACCUMULATE onto an explicit review
session, append-only, supersession VISIBLE, per-source coverage attribution, cross-slice
decisions do NOT carry (client_id provenance-only)" invariant is pinned in the panel spec, will
be enforced in agent/review_store.py + api/app.py, and is asserted here.

GAP REPORTED (honest status -- deviation from the task premise): the task said "today
agent/review_store.py doesn't exist and the endpoints 404/405". In THIS worktree the STORE
module already exists and is functional (create_review / append_slice / merged_view /
list_reviews) but is COMPLETELY UNWIRED -- api/app.py has no /review-session routes and
post_review_upload has no review_id Form field. So the failing-first reason for every test below
is the MISSING API SURFACE / MISSING UPLOAD WIRING, not a missing module. Because the store
exists, tests that need a valid review_id mint one directly via ``review_store.create_review()``
so the upload path actually RUNS -- isolating the wiring gap. NOTE also: the existing
``create_review()`` returns the create RECORD dict (``{"review_id", ...}``), NOT the bare id
string the task text described ("-> generated id like r_<hex>"); the tests pin the ACTUAL module
contract (``["review_id"]``) to avoid a wrong-reason failure, and this discrepancy is flagged.

HERMETIC: tmp_path everything; engine PDF dir + audit root + decision store + review store all
redirected to tmp. The per-upload fresh-audit-subdir trick (seal.run_ts has seconds resolution ->
two uploads in one second collide on one read-only bundle dir; sidestepped per-call, NEVER patched
in the engine). xero_demo / xero_sales_demo / extract_demo are file-import (non-SAP) clients loaded
with check_connectivity=False; dummy SAP creds are set (harmless) mirroring the decision-reapply
fixture. Pure stdlib + pytest + fastapi.testclient + openpyxl (T5 only). No anthropic import.
SAP off, no tokens, no network.
"""
from __future__ import annotations

import itertools
import json
import re
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

import audit_bundle.seal as _seal
from api.app import app

# Each upload gets its OWN audit subdir (seal.run_ts seconds resolution -> per-call bump).
_UPLOAD_SEQ = itertools.count()

_REPO_ROOT = Path(__file__).resolve().parents[1]

_F5_FIXTURE = (
    _REPO_ROOT / "tests" / "fixtures" / "xero-f5-export"
    / "AgentAssist_IRAS_F5_2026-04-01_to_2026-06-30.xlsx"
)
_F5_NAME = "AgentAssist_IRAS_F5_2026-04-01_to_2026-06-30.xlsx"

_SALES_FIXTURE = (
    _REPO_ROOT / "tests" / "fixtures" / "xero-sales-export"
    / "AgentAssist_Xero_SalesInvoices_2026-04-01_to_2026-06-30.xlsx"
)
_SALES_NAME = "AgentAssist_Xero_SalesInvoices_2026-04-01_to_2026-06-30.xlsx"

# The 5-key top-level F5 upload literal (R1 -- must NEVER gain a review_id key).
_F5_TOP_KEYS = {"source_kind", "validation_status", "disclaimer", "coverage_status", "queue"}

# The review_id char guard mirrored from agent.review_store._REVIEW_ID_RE.
_REVIEW_ID_RE = re.compile(r"^[a-z0-9_]+$")


@pytest.fixture()
def client() -> TestClient:
    return TestClient(app)


@pytest.fixture()
def hermetic(tmp_path, monkeypatch):
    """Engine PDF dir + audit root + decision store + review store all redirected to tmp.

    The review-store redirect is what keeps the accumulation store hermetic: the (existing)
    ``agent.review_store._REVIEWS_ROOT`` module constant is repointed at ``tmp_path/reviews``
    so both ``create_review()`` (test-side minting) and the Phase-2 upload wiring write there.
    """
    monkeypatch.setattr("engine.review._REPORTS_DIR", tmp_path / "reports")
    monkeypatch.setattr("audit_bundle.seal._AUDIT_ROOT", tmp_path / "audit")
    monkeypatch.setattr("agent.decision_store._DECISIONS_ROOT", tmp_path / "decisions")
    monkeypatch.setenv("SAP_USERNAME", "dummy")
    monkeypatch.setenv("SAP_PASSWORD", "dummy")
    # Import inside the fixture so a future removal of the module cannot break collection;
    # raising=False keeps the redirect resilient to the constant being renamed.
    import agent.review_store as _rs
    monkeypatch.setattr(_rs, "_REVIEWS_ROOT", tmp_path / "reviews", raising=False)
    return tmp_path


def _bump_audit() -> None:
    _seal._AUDIT_ROOT = _seal._AUDIT_ROOT.parent / f"audit-{next(_UPLOAD_SEQ)}"


def _upload_resp(client: TestClient, file_bytes: bytes, name: str, review_id=None):
    """POST /review/upload with a fresh audit subdir; returns the raw Response (no assert)."""
    _bump_audit()
    data = {"review_id": review_id} if review_id is not None else None
    return client.post("/review/upload", files={"file": (name, file_bytes)}, data=data)


def _post_upload(client: TestClient, file_bytes: bytes, name: str, review_id=None) -> dict:
    resp = _upload_resp(client, file_bytes, name, review_id=review_id)
    assert resp.status_code == 200, resp.text
    return resp.json()


def _new_review(label: str = "") -> str:
    """Mint a valid review_id directly via the (existing) store -- the upload wiring is the
    thing under test, so we set up a real session without depending on POST /review-session."""
    import agent.review_store as rs
    return rs.create_review(label)["review_id"]


def _merged(client: TestClient, rid: str):
    return client.get(f"/review-session/{rid}")


def _f5_slice(view: dict) -> dict:
    return next(s for s in view["slices"] if s["source_kind"] == "xero_f5_upload")


def _sales_slice(view: dict) -> dict:
    return next(s for s in view["slices"] if s["source_kind"] == "xero_sales_upload")


# -- T1 -- create + get + list session lifecycle over the new endpoints ----------------------

def test_t1_create_get_list_session(client, hermetic):
    """POST /review-session creates a session; GET reads it back empty; the list contains it.

    FAILS TODAY: the /review-session routes do not exist -- POST returns 404 (route absent), so
    the very first assertion (200 on create) is red. This pins the whole session-lifecycle
    surface the accumulation feature is built on.
    """
    assert _F5_FIXTURE.is_file(), f"committed Xero F5 fixture missing: {_F5_FIXTURE}"

    created = client.post("/review-session", json={})
    assert created.status_code == 200, created.text
    body = created.json()
    rid = body["review_id"]
    assert _REVIEW_ID_RE.fullmatch(rid), f"review_id must match ^[a-z0-9_]+$, got {rid!r}"
    assert body["created_at"], "create response must carry created_at"

    got = _merged(client, rid)
    assert got.status_code == 200, got.text
    assert got.json()["slices"] == [], "a fresh session has no slices"

    listing = client.get("/review-session")
    assert listing.status_code == 200, listing.text
    assert any(r["review_id"] == rid for r in listing.json()["reviews"]), (
        "the created session must appear in the list view"
    )

    assert client.get("/review-session/nonexistent_id").status_code == 404, (
        "an unknown (but valid-shaped) session id is a 404 (GET /review 404-shape parity)"
    )


# -- T2 -- R1 byte-identical pin: review_id NEVER changes the upload response --------------

def test_t2_r1_review_id_is_byte_identical(client, hermetic):
    """R1: a valid review_id changes ONLY the store side-effect, never the upload response.

    Two F5 uploads over the identical fixture -- one WITHOUT review_id (control), one WITH a
    valid review_id -- must return EXACTLY the same 5-key body (deterministic over one fixture:
    empty decision store -> demoted False, deterministic dossier framing/inputs_hash), and the
    with-id body must carry NO review_id key.

    FAILS TODAY: today the review_id multipart part is IGNORED by FastAPI (the endpoint has no
    such parameter), so the two bodies ARE already identical -- that half would pass vacuously.
    The pin that BITES is the slice attachment: after the with-id upload the session must show
    slice_count 1. GET /review-session/{id} is a 404 today (route absent) AND no slice was
    written (upload wiring absent) -> that assertion is the red one, and it stays meaningful
    after the Form field is added (a no-op wiring that accepts-but-drops review_id is caught).
    """
    rid = _new_review()  # store exists -> mint a real session so both uploads run
    f5 = _F5_FIXTURE.read_bytes()

    control = _post_upload(client, f5, _F5_NAME)
    with_id = _post_upload(client, f5, _F5_NAME, review_id=rid)

    # R1 -- key set is the 5-key literal on both; no review_id key ever leaks into the response.
    assert set(control.keys()) == _F5_TOP_KEYS
    assert set(with_id.keys()) == _F5_TOP_KEYS
    assert "review_id" not in with_id
    for key in ("source_kind", "validation_status", "disclaimer"):
        assert with_id[key] == control[key]
    # Deterministic over the same fixture + empty store -> the whole bodies are equal.
    assert with_id == control, "R1: review_id must not alter the upload response body"

    # The pin that bites: the with-id upload attached exactly one slice to the session.
    view = _merged(client, rid)
    assert view.status_code == 200, view.text
    assert len(view.json()["slices"]) == 1, "a valid review_id must attach the upload as a slice"


# -- T3 -- accumulation across export types (F5 + sales) into one session --------------------

def test_t3_accumulation_across_export_types(client, hermetic):
    """Two DIFFERENT export types accumulate as TWO slices under one session (R6 grouped view).

    Upload the F5 fixture and the committed sales fixture with the SAME review_id, then read the
    merged view: two active slices (source_kinds {xero_f5_upload, xero_sales_upload}); the F5
    slice's stored queue carries the E2/E3/E4 detect rows WITH populated candidate_framing_text
    (the post-dossier serialized rows); the sales slice queue carries NO E-check rows -- but it
    is NOT empty: the clean sales-only fixture legitimately fires the COMPLETENESS volume check
    ("Purchase volume below completeness threshold" -- zero purchase invoices on a sales-only
    export), corrected here from the original "queue is []" draft after running against
    reality; coverage_matrix carries (a) every check from both slices and (b) PER-SOURCE
    attribution -- NO_GST_REG has TWO entries, one per source_kind, each level "unavailable" with
    a non-empty reason.

    FAILS TODAY: GET /review-session/{id} is a 404 (route absent) and the uploads write no
    slices (wiring absent) -> the merged-view read is red.
    """
    assert _SALES_FIXTURE.is_file(), f"committed Xero sales fixture missing: {_SALES_FIXTURE}"
    rid = _new_review("accumulation")

    _post_upload(client, _F5_FIXTURE.read_bytes(), _F5_NAME, review_id=rid)
    _post_upload(client, _SALES_FIXTURE.read_bytes(), _SALES_NAME, review_id=rid)

    resp = _merged(client, rid)
    assert resp.status_code == 200, resp.text
    view = resp.json()

    assert view["validation_status"] == "unvalidated"
    assert isinstance(view["disclaimer"], str) and view["disclaimer"].strip()

    kinds = {s["source_kind"] for s in view["slices"]}
    assert kinds == {"xero_f5_upload", "xero_sales_upload"}, (
        "both export types accumulate as distinct active slices"
    )

    f5 = _f5_slice(view)
    detect = [r for r in f5["queue"] if str(r.get("finding_id", "")).startswith("detect:")]
    assert detect, "the F5 slice must retain its detect rows"
    codes = {r["check_id"] for r in detect}
    assert {"E2", "E3", "E4"} <= codes, "F5 fixture yields E2/E3/E4"
    for row in detect:
        if row["check_id"] in {"E2", "E3", "E4"}:
            assert row["candidate_framing_text"].strip(), (
                "stored F5 detect rows carry populated dossier framing (serialized rows)"
            )

    sales = _sales_slice(view)
    sales_codes = {r.get("check_id") for r in sales["queue"]}
    assert not (sales_codes & {"E1", "E2", "E3", "E4"}), (
        "the clean sales fixture yields no E-check findings"
    )
    # The sales-only export legitimately fires the COMPLETENESS volume check (zero
    # purchase invoices) -- an honest finding, retained in the slice, not scrubbed.
    assert sales_codes <= {"COMPLETENESS"}, (
        f"unexpected sales-slice finding types: {sales_codes}"
    )

    matrix = view["coverage_matrix"]
    f5_checks = {r["check"] for r in f5["coverage_status"]}
    sales_checks = {r["check"] for r in sales["coverage_status"]}
    assert set(matrix.keys()) >= (f5_checks | sales_checks), (
        "coverage_matrix carries every check name from both slices"
    )
    nogst = matrix["NO_GST_REG"]
    assert {e["source_kind"] for e in nogst} == {"xero_f5_upload", "xero_sales_upload"}, (
        "per-source attribution: one NO_GST_REG entry per source_kind"
    )
    for entry in nogst:
        assert entry["level"] == "unavailable"
        assert isinstance(entry["reason"], str) and entry["reason"].strip()


# -- T4 -- idempotent re-upload: same bytes twice is one slice ------------------------------

def test_t4_idempotent_reupload(client, hermetic):
    """R2 idempotency: re-uploading the SAME F5 bytes into one session does NOT add a slice.

    Two identical uploads -> slice_count stays 1, and the on-disk JSONL has the create line plus
    exactly ONE slice line (read directly under the monkeypatched _REVIEWS_ROOT).

    FAILS TODAY: the upload ignores review_id (wiring absent) so NO slice line is ever written --
    the store file has only the create line and slice_count is unreadable (list route 404). The
    JSONL read (expecting create + 1 slice) is the red assertion.
    """
    import agent.review_store as rs

    rid = _new_review()
    f5 = _F5_FIXTURE.read_bytes()
    _post_upload(client, f5, _F5_NAME, review_id=rid)
    _post_upload(client, f5, _F5_NAME, review_id=rid)  # identical -> idempotent skip

    listing = client.get("/review-session")
    assert listing.status_code == 200, listing.text
    row = next(r for r in listing.json()["reviews"] if r["review_id"] == rid)
    assert row["slice_count"] == 1, "same (source_kind, sha256) must not add a second slice"

    path = rs._REVIEWS_ROOT / rid / "slices.jsonl"
    lines = [json.loads(ln) for ln in path.read_text(encoding="utf-8").splitlines() if ln.strip()]
    creates = [r for r in lines if r.get("type") == "create"]
    slices = [r for r in lines if r.get("type") == "slice"]
    assert len(creates) == 1, "exactly one create record"
    assert len(slices) == 1, "idempotent re-upload keeps exactly one slice line"


# -- T5 -- supersede-in-view: different bytes, same source_kind -----------------------------

def test_t5_supersede_in_view(client, hermetic):
    """R2 supersession: a DIFFERENT-sha F5 supersedes the old one VISIBLY, append-only on disk.

    Upload the F5 fixture, then a MODIFIED F5 (a genuinely different sha256, still a valid F5 the
    engine completes over) with the same review_id. Merged view: ONE active xero_f5_upload slice
    (the new sha), and the OLD sha in "superseded" with superseded_at set. The store JSONL keeps
    BOTH slice lines (append-only -- superseded records are RETAINED).

    FAILS TODAY: the uploads write no slices (wiring absent) and GET /review-session/{id} is a
    404 -> the supersession view + the two-slice-line JSONL are both red.
    """
    import agent.review_store as rs

    rid = _new_review()
    _post_upload(client, _F5_FIXTURE.read_bytes(), _F5_NAME, review_id=rid)

    v1 = _merged(client, rid)
    assert v1.status_code == 200, v1.text
    old_sha = _f5_slice(v1.json())["sha256"]

    modified = _modified_f5_bytes(hermetic)
    _post_upload(client, modified, _F5_NAME, review_id=rid)

    v2 = _merged(client, rid)
    assert v2.status_code == 200, v2.text
    view = v2.json()

    active = [s for s in view["slices"] if s["source_kind"] == "xero_f5_upload"]
    assert len(active) == 1, "supersession keeps exactly one ACTIVE slice per source_kind"
    assert active[0]["sha256"] != old_sha, "the active slice is the NEW sha"

    superseded = view["superseded"]
    assert any(
        x["sha256"] == old_sha and x.get("superseded_at")
        for x in superseded
    ), "the old sha must appear in superseded with superseded_at set (supersession VISIBLE)"

    path = rs._REVIEWS_ROOT / rid / "slices.jsonl"
    lines = [json.loads(ln) for ln in path.read_text(encoding="utf-8").splitlines() if ln.strip()]
    slices = [r for r in lines if r.get("type") == "slice"]
    assert len(slices) == 2, "append-only: BOTH slice lines retained on disk"
    assert {s.get("sha256") for s in slices} == {old_sha, active[0]["sha256"]}


def _modified_f5_bytes(dst_dir: Path) -> bytes:
    """A copy of the F5 fixture with one cell changed so its sha256 differs, engine still runs.

    Bumps a Description cell (a column the value-box loader IGNORES) so the review output is
    unchanged while the file bytes differ; falls back to a spare marker sheet if the Description
    column cannot be located. openpyxl round-trip is safe here -- a Xero export is a values-only
    workbook (no formulas to lose).
    """
    from openpyxl import load_workbook
    from feeders.xero_f5_reader import (
        TRANSACTIONS_SHEET,
        _COL_DATE,
        _COL_DESCRIPTION,
        _COL_TAX_RATE,
    )

    wb = load_workbook(_F5_FIXTURE)  # default: not read_only, keeps values
    ws = wb[TRANSACTIONS_SHEET]

    header_row = None
    desc_col = None
    for row in ws.iter_rows(min_row=1, max_row=12):
        labels = [str(c.value).strip() if c.value is not None else "" for c in row]
        if _COL_DATE in labels and _COL_TAX_RATE in labels:
            header_row = row[0].row
            if _COL_DESCRIPTION in labels:
                desc_col = labels.index(_COL_DESCRIPTION) + 1  # 1-based column index
            break

    changed = False
    if header_row is not None and desc_col is not None:
        for r in range(header_row + 1, ws.max_row + 1):
            cell = ws.cell(row=r, column=desc_col)
            if cell.value not in (None, ""):
                cell.value = f"{cell.value} [edited-for-supersede]"
                changed = True
                break
    if not changed:
        marker = wb.create_sheet("__supersede_marker__")
        marker["A1"] = "supersede-marker"

    out = dst_dir / "modified_f5.xlsx"
    wb.save(out)
    return out.read_bytes()


# -- T6 -- decisions still re-apply inside a review (B3a non-regression) ---------------------

def test_t6_decisions_reapply_inside_review(client, hermetic):
    """A persisted decision RE-APPLIES onto a slice stored under a session (decisions ride rows).

    Upload F5 WITHOUT review_id first to read a finding's fingerprint; POST /decision (Mark known)
    under client_id xero_demo; THEN upload the SAME F5 WITH review_id. The stored slice's queue
    rows must show the matching finding demoted True with KNOWN_ACCEPTED -- re-application runs
    BEFORE storage, so the durable slice carries the demotion (present-but-demoted, never dropped).

    (The same bytes are re-used deliberately -- decisions key on the finding FINGERPRINT, not the
    file sha; the first upload is review_id-less so it does not trigger the same-sha idempotent
    skip on the second, review_id-bearing upload.)

    FAILS TODAY: GET /review-session/{id} is a 404 and no slice is written (wiring absent) -> the
    stored-row read is red. POST /decision + the plain upload already work (B3a is merged).
    """
    f5 = _F5_FIXTURE.read_bytes()
    base = _post_upload(client, f5, _F5_NAME)  # no review_id -> stateless, gives a fingerprint
    target = next(r for r in base["queue"] if r["error_code"] == "E4")
    fp = target["fingerprint"]
    assert isinstance(fp, str) and fp.startswith("sha256:")

    decided = client.post("/decision", json={
        "client_id": "xero_demo",
        "finding_id": target["finding_id"],
        "fingerprint": fp,
        "action": "Mark known",
        "note": "Standing accepted treatment for this supplier.",
        "reviewer_name": "Collin",
    })
    assert decided.status_code == 200, decided.text

    rid = _new_review()
    _post_upload(client, f5, _F5_NAME, review_id=rid)

    resp = _merged(client, rid)
    assert resp.status_code == 200, resp.text
    f5 = _f5_slice(resp.json())
    matched = [r for r in f5["queue"] if r["fingerprint"] == fp]
    assert matched, "the marked-known finding must be present in the stored slice"
    for row in matched:
        assert row["demoted"] is True, "a persisted Mark known demotes the stored row"
        assert "KNOWN_ACCEPTED" in (row["prior_dispositions"] or [])


# -- T7 -- store hygiene: monkeypatchable root + path/format guards --------------------------

def test_t7_store_hygiene_and_guards(client, hermetic):
    """The store root is monkeypatchable and the upload validates the review_id path segment.

    Store level (already works -- the module exists): _REVIEWS_ROOT is redirected under tmp and
    create_review writes reviews/<id>/slices.jsonl there. Endpoint level (Phase 2): a review_id
    with invalid chars ("../evil" / "UPPER") -> 422 with NO store write; a valid-shaped but
    unknown id -> 404.

    FAILS TODAY: the upload ignores review_id (no such Form field), so the invalid-char upload
    returns 200 instead of 422 and the unknown-id upload returns 200 instead of 404 -- those are
    the red assertions. (The create_review file-write half passes today; the guard half is the
    Phase-2 behaviour being pinned.)
    """
    import agent.review_store as rs

    reviews_root = hermetic / "reviews"
    monkeypatch_ok = rs._REVIEWS_ROOT == reviews_root
    assert monkeypatch_ok, "the hermetic fixture must have redirected _REVIEWS_ROOT under tmp"

    rec = rs.create_review()
    rid = rec["review_id"]
    assert (rs._REVIEWS_ROOT / rid / "slices.jsonl").is_file(), (
        "create_review writes reviews/<id>/slices.jsonl under the monkeypatched root"
    )
    before = {p.name for p in rs._REVIEWS_ROOT.iterdir() if p.is_dir()}

    f5 = _F5_FIXTURE.read_bytes()
    for bad in ("../evil", "UPPER"):
        resp = _upload_resp(client, f5, _F5_NAME, review_id=bad)
        assert resp.status_code == 422, (
            f"invalid review_id {bad!r} must be a 422, got {resp.status_code}"
        )

    unknown = _upload_resp(client, f5, _F5_NAME, review_id="r_doesnotexist9")
    assert unknown.status_code == 404, (
        "a valid-shaped but unknown review_id must be a 404 (never a silent new session)"
    )

    after = {p.name for p in rs._REVIEWS_ROOT.iterdir() if p.is_dir()}
    assert after == before, "an invalid/unknown review_id must NOT write to the store"


# -- T8 -- cross-slice decision NON-carry, pinned honestly (R5) ------------------------------

def test_t8_cross_slice_decision_non_carry_structural(client, hermetic):
    """Cross-slice decisions do NOT carry: slices carry DIFFERENT client_ids in one session.

    HONEST DRIVE (R5): a decision recorded under client_id ``xero_demo`` (the F5 slice) must not
    demote a same-fingerprint row on the ``xero_sales_demo`` sales slice -- they read DIFFERENT
    decision stores (decision_store keys on client_id, review_store keys on review_id, and the
    slice records the client_id as PROVENANCE ONLY). The committed sales fixture yields no
    E-CHECK findings (its only row is the COMPLETENESS purchase-volume check, whose fingerprint
    counterparty differs from any F5 vendor row), so there is no same-fingerprint pair to
    exercise the non-carry directly; the STRUCTURAL fact behind it is asserted instead -- the
    two slices in ONE session carry DISTINCT client_id values (xero_demo vs xero_sales_demo). UX consequence (the reason this is pinned):
    adjudicating a finding on the F5 slice does NOT ripple to the sales slice; unifying
    adjudication across an entity's slices is #45 territory, named per Terry R5.

    FAILS TODAY: GET /review-session/{id} is a 404 and no slices are written (wiring absent) ->
    the client_id-provenance read is red.
    """
    rid = _new_review("cross-slice")
    _post_upload(client, _F5_FIXTURE.read_bytes(), _F5_NAME, review_id=rid)
    _post_upload(client, _SALES_FIXTURE.read_bytes(), _SALES_NAME, review_id=rid)

    resp = _merged(client, rid)
    assert resp.status_code == 200, resp.text
    view = resp.json()

    f5_client = _f5_slice(view)["client_id"]
    sales_client = _sales_slice(view)["client_id"]
    assert f5_client == "xero_demo", "F5 slice records its config client_id (provenance only)"
    assert sales_client == "xero_sales_demo", "sales slice records its own config client_id"
    assert f5_client != sales_client, (
        "distinct per-slice client_ids: a decision under one does NOT carry to the other"
    )
