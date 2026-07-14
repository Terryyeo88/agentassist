"""tests/test_extract_engine_upload.py — BUILD 3 failing-first tests: turn the general-extract
engine path ON (default) under a DEMO config, gated by a global kill switch.

Written BEFORE the implementation; MUST fail today for the RIGHT reason — the non-Xero .xlsx
branch is coverage-only (engine not run), there is no ``extract_review`` source_kind, no
``config_scope`` marker, no ``AGENTASSIST_EXTRACT_ENGINE`` flag, and no ``extract_demo`` config.

WHAT THIS LOCKS:
  * BT1 — a generic (non-Xero) extract .xlsx, flag ON, runs the engine under the demo config
    and returns findings in the SHARED ``queue`` shape; validation_status stays "unvalidated";
    no ai_candidates leak.
  * BT2 — the kill switch is GLOBAL and DEFAULT-ON: unset → engine (findings); explicit OFF →
    the pre-build coverage-only body. No per-client flag exists on this path.
  * BT3 — off-format degrades / errors, never GUESSES: a non-extract workbook or a non-numeric
    value in a numeric column → 422, and NO findings. (Recast T3 regression — the reader is
    fixed-schema and cannot infer columns; turning the engine on introduces no guessed finding.)
  * BT4 — demo-config honesty: the body carries a ``config_scope="default_demo"`` marker so the
    surface can TELL it ran under the demo config, not the uploader's.

HONEST STATUS: real findings over SYNTHETIC data under a DEFAULT DEMO config — CANDIDATES, never
verdicts; validation_status "unvalidated"; synthetic-format at best; T2.11 unmoved.

Pure stdlib + pytest + fastapi.testclient + openpyxl. No anthropic. SAP off, hermetic.
"""
from __future__ import annotations

import importlib.util
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from api.app import app
from api.viewmodel import QUEUE_ITEM_KEYS

_REPO_ROOT = Path(__file__).resolve().parents[1]
_FROZEN_EXTRACT_DIR = _REPO_ROOT / "tests" / "fixtures" / "sbodemosg-extract"
_FLAG = "AGENTASSIST_EXTRACT_ENGINE"

# synth exporter — local importlib load (tests/ is not a package; do NOT import it as a module).
_synth_spec = importlib.util.spec_from_file_location(
    "extract_engine_synth_export", Path(__file__).resolve().parent / "synth_extract_export.py"
)
_synth = importlib.util.module_from_spec(_synth_spec)
_synth_spec.loader.exec_module(_synth)


@pytest.fixture()
def client() -> TestClient:
    return TestClient(app)


@pytest.fixture()
def hermetic_engine(tmp_path, monkeypatch):
    """Redirect the engine's PDF dir + the bundle's audit root to tmp, and set dummy SAP creds.
    Mirrors test_xero_engine_upload.py — the extract engine path runs the FULL review()."""
    monkeypatch.setattr("engine.review._REPORTS_DIR", tmp_path / "reports")
    monkeypatch.setattr("audit_bundle.seal._AUDIT_ROOT", tmp_path / "audit")
    monkeypatch.setenv("SAP_USERNAME", "dummy")
    monkeypatch.setenv("SAP_PASSWORD", "dummy")


def _synthetic_extract_bytes(tmp_path: Path) -> bytes:
    out = tmp_path / "synthetic_export.xlsx"
    _synth.export_xlsx(_FROZEN_EXTRACT_DIR, out)
    return out.read_bytes()


def _post(client: TestClient, content: bytes, filename: str = "export.xlsx"):
    return client.post("/review/upload", files={"file": (filename, content)})


# ── BT1 — path ON (default): engine runs under demo config, returns queue ──

def test_extract_upload_runs_engine_under_demo_config(client, hermetic_engine, tmp_path, monkeypatch):
    monkeypatch.delenv(_FLAG, raising=False)  # default = ON
    resp = _post(client, _synthetic_extract_bytes(tmp_path))
    assert resp.status_code == 200, resp.text
    body = resp.json()

    assert body["source_kind"] == "extract_review"
    assert "queue" in body and isinstance(body["queue"], list) and body["queue"]
    for item in body["queue"]:
        assert set(item.keys()) == set(QUEUE_ITEM_KEYS)
        assert item["validation_status"] == "unvalidated"
    assert body["validation_status"] == "unvalidated"
    assert "ai_candidates" not in body
    # coverage_status still surfaced alongside the findings.
    assert isinstance(body["coverage_status"], list) and body["coverage_status"]


# ── BT2 — global kill switch, DEFAULT-ON ──

def test_kill_switch_is_global_and_default_on(client, hermetic_engine, tmp_path, monkeypatch):
    content = _synthetic_extract_bytes(tmp_path)

    # OFF (explicit) → pre-build coverage-only body, engine NOT run.
    monkeypatch.setenv(_FLAG, "off")
    off = _post(client, content).json()
    assert off["source_kind"] == "extract_upload"
    assert "queue" not in off
    assert set(off.keys()) == {"source_kind", "validation_status", "disclaimer", "coverage_status"}

    # UNSET → default ON → engine runs (findings).
    monkeypatch.delenv(_FLAG, raising=False)
    on = _post(client, content).json()
    assert on["source_kind"] == "extract_review"
    assert on["queue"]


# ── BT3 — off-format degrades/errors, never GUESSES a finding ──

def test_offformat_never_guesses_a_finding(client, hermetic_engine, tmp_path, monkeypatch):
    from openpyxl import Workbook

    monkeypatch.delenv(_FLAG, raising=False)  # ON — prove the engine-on path still 422s off-format

    # (a) a generic .xlsx that is NOT the extract format (no canonical sheets) → 422, no findings.
    wb = Workbook()
    wb.active.title = "RandomSheet"
    wb.active.append(["whatever", "columns"])
    wb.active.append(["a", "b"])
    generic = tmp_path / "generic.xlsx"
    wb.save(generic)
    r1 = _post(client, generic.read_bytes())
    assert r1.status_code == 422, r1.text
    assert "queue" not in r1.json()

    # (b) canonical sheets but a NON-NUMERIC value in a numeric column → 422 (never a guess).
    wb2 = Workbook()
    ws_d = wb2.active
    ws_d.title = "documents"
    ws_d.append(["doc_type", "doc_kind", "DocNum", "DocDate", "CardCode", "CardName",
                 "DocCurrency", "DocTotal", "line_index", "VatGroup", "LineTotal", "TaxTotal"])
    ws_d.append(["sales", "invoice", "1", "2024-07-01", "C1", "Acme",
                 "SGD", "100", "0", "SR", "NOT_A_NUMBER", "9"])
    wb2.create_sheet("business_partners").append(["CardCode", "CardName", "FederalTaxID"])
    wb2.create_sheet("listing").append(["scope", "doc_type", "DocNum", "Series", "Cancelled",
                                        "CardCode", "NumAtCard", "DocTotal"])
    badnum = tmp_path / "badnum.xlsx"
    wb2.save(badnum)
    r2 = _post(client, badnum.read_bytes())
    assert r2.status_code == 422, r2.text
    assert "queue" not in r2.json()


# ── BT4 — demo-config honesty marker ──

def test_body_carries_default_config_marker(client, hermetic_engine, tmp_path, monkeypatch):
    monkeypatch.delenv(_FLAG, raising=False)
    body = _post(client, _synthetic_extract_bytes(tmp_path)).json()
    assert body["config_scope"] == "default_demo", (
        "the surface must be able to tell it ran under the demo config, not the uploader's"
    )
