"""T2.24 PR-2 — feeders.xero_f5_reader.parse_declared_return (FAILING-FIRST).

Pins the product parser that reads Box 6 / Box 7 VALUES off the F5 workbook's 'Return'
sheet by STRUCTURAL label scan (col A == "Box 6"/"Box 7", value in col C / index 2) — no
fixed header offset. The returned keys are EXACTLY {output_tax, input_tax} so the dict
drops straight into the gst_ledger["declared_boxes"] seam with no mapping shim.

This RETIRES the test-scaffolding _declared_boxes_from_return in
test_t224_ledger_recon_internal_consistency.py — parse_declared_return is the product path.

Real-FORMAT over SYNTHETIC content; NOT accuracy-validated (T2.11 unmoved).
"""
from __future__ import annotations

from pathlib import Path

import pytest

from feeders.xero_f5_reader import parse_declared_return

_FIX = Path(__file__).resolve().parent / "fixtures" / "xero-real-format"
_F5_WORKBOOK = _FIX / "AgentAssist_IRAS_F5_2026-04-01_to_2026-06-30.xlsx"
# The GST control-account ledger has NO 'Return' sheet — the honest-failure input.
_LEDGER_WORKBOOK = _FIX / "AgentAssist_-_Account_Transactions.xlsx"


def test_parse_declared_return_reads_box6_box7_values():
    """Box 6 = 720.0 (output tax), Box 7 = 1271.3 (input tax); keys EXACTLY those two."""
    boxes = parse_declared_return(_F5_WORKBOOK)
    assert boxes == {"output_tax": 720.0, "input_tax": 1271.3}
    # Exact key contract — no extra keys, no missing keys (drops into the seam unmapped).
    assert set(boxes) == {"output_tax", "input_tax"}
    assert isinstance(boxes["output_tax"], float)
    assert isinstance(boxes["input_tax"], float)


def test_missing_return_sheet_raises_valueerror():
    """A workbook with no 'Return' sheet fails LOUD (ValueError), never a silent None."""
    with pytest.raises(ValueError):
        parse_declared_return(_LEDGER_WORKBOOK)


def test_return_sheet_missing_box6_row_raises_valueerror(tmp_path):
    """A 'Return' sheet that HAS the sheet but is missing the Box 6 row raises ValueError.

    Built in tmp_path so the failure is on structural absence of the Box 6 label row, not
    a bad sheet name — the parser must not silently return None / a partial dict.
    """
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Return"
    # Only Box 7 present; Box 6 row deliberately absent.
    ws["A19"] = "Box 7"
    ws["C19"] = 1271.3
    broken = tmp_path / "F5_missing_box6.xlsx"
    wb.save(broken)

    with pytest.raises(ValueError):
        parse_declared_return(broken)
