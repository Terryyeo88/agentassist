"""feeders/extract_reader.py — ExtractChainReader: a ChainReader over a client GST export.

Reads a client Excel/CSV GST export (a directory of CSVs, or a single .xlsx workbook
with the same sheet names) and emits the SAME shaped record lists/dicts the deterministic
checking core consumes from the live-SAP feeder. Satisfies the ``ChainReader`` contract
(T2.23) STRUCTURALLY — it neither imports nor widens the Protocol:

    count(entity, period_start, period_end)            -> Optional[int]
    fetch_invoices(entity, period_start, period_end)   -> list[dict]
    fetch_credit_notes(entity_type, ps, pe)            -> list[dict]
    get_business_partner(card_code)                    -> dict
    fetch_listing(period)                              -> dict   (4 buckets)

Normalisation happens AT THE FEEDER (export columns → canonical fields); the checking
core is untouched and tax-code normalisation still flows through its
``normalize_vat_group`` (T2.19). One machine, two feeders.

THE COUNT TRAP (backlog #5): a client export contains every row, so ``count()`` returns
the TRUE record count — unlike the live ``SapChainReader.count`` / frozen oracle, which
reproduce today's dormant ``@odata.count`` → None. ``count()`` here is therefore unit-tested
against the export's known row count and is NEVER asserted against the frozen S0 / oracle.

Beside the Protocol methods, ``coverage()`` declares which canonical fields the loaded
export carried (header AND value population), and ``coverage_status()`` maps that onto the
per-check data-coverage status the working paper surfaces (T2.12 slice 2B).

Pure stdlib; openpyxl is imported lazily only on the .xlsx path. No SAP, no anthropic.
"""
from __future__ import annotations

import copy
import csv
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

from feeders import extract_schema as schema
from feeders.coverage_status import derive_coverage_statuses


@dataclass(frozen=True)
class ExtractCoverage:
    """Per-(surface, field) coverage declaration for one loaded export.

    ``fields`` maps each ``(surface, field)`` key (see ``schema.COVERAGE_FIELDS``) to whether
    the export carried its source COLUMN (header presence). ``populated`` maps the same keys
    to whether that column carried at least one NON-EMPTY value. The two differ when a column
    is present but 0%-populated (e.g. ``NumAtCard`` in SBODEMOSG) — header-present yet value-
    empty. ``is_covered()`` is the value-aware test slice 2B's status mapping uses; ``is_full()``
    stays HEADER-fullness (all columns present) for backward compatibility.

    EMISSION seam only — ``feeders.coverage_status`` maps this onto per-check status; it does
    not live on the ChainReader Protocol.
    """

    fields: dict       # (surface, field) -> bool   column header present
    populated: dict    # (surface, field) -> bool   column carried >=1 non-empty value

    def is_full(self) -> bool:
        """All source columns present (header-fullness — NOT value-population-aware)."""
        return all(self.fields.values())

    def missing(self) -> list:
        """The (surface, field) keys whose source column is absent from the export."""
        return sorted(k for k, present in self.fields.items() if not present)

    def is_present(self, surface, field) -> bool:
        """Whether the column for ``(surface, field)`` appeared in the export header."""
        return bool(self.fields.get((surface, field), False))

    def is_populated(self, surface, field) -> bool:
        """Whether the column for ``(surface, field)`` carried >=1 non-empty value."""
        return bool(self.populated.get((surface, field), False))

    def is_covered(self, surface, field) -> bool:
        """Value-aware coverage: the column is BOTH present AND populated.

        A present-but-0%-populated column is NOT covered — this is the distinction
        that keeps a silently-partial check (DUP_CLAIM under-detection, NO_GST_REG
        under a blank supplier master) from reading as ``full``.
        """
        return self.is_present(surface, field) and self.is_populated(surface, field)


class ExtractChainReader:
    """A ``ChainReader`` backed by a client Excel/CSV GST export.

    Args:
        source: a directory containing ``documents.csv`` / ``business_partners.csv`` /
                ``listing.csv``, OR a single ``.xlsx`` workbook with sheets of the same
                base names.

    Each read returns FRESH objects (deepcopy on hand-off) so the ``is_credit_note`` tag
    on the credit-note read can never leak into the untagged invoice read via a shared
    mutable — the per-call-freshness hazard the live feeder avoids naturally.
    """

    def __init__(self, source):
        self._source = Path(source)
        sheets, present_columns = _load_sheets(self._source)
        self._present_columns = present_columns
        # Value-population per (sheet → set of columns that carried >=1 non-empty cell).
        # Computed from the loaded rows before they are projected away, so coverage can
        # distinguish a present-but-0%-populated column from a populated one.
        self._populated_columns = _compute_populated_columns(sheets)

        # --- documents → grouped canonical documents, bucketed by (doc_type, doc_kind) ---
        self._docs_by_bucket: dict[tuple, list[dict]] = {
            ("sales", "invoice"): [],
            ("purchase", "invoice"): [],
            ("sales", "credit_note"): [],
            ("purchase", "credit_note"): [],
        }
        self._build_documents(sheets[schema.DOCUMENTS_SHEET])

        # --- business partners → canonical dict keyed by CardCode ---
        self._bp_by_card: dict[str, dict] = {}
        for row in sheets[schema.BUSINESS_PARTNERS_SHEET]:
            card_code = schema.to_str(row.get("CardCode"))
            self._bp_by_card[card_code] = {
                "CardCode": card_code,
                "FederalTaxID": schema.to_opt_str(row.get("FederalTaxID")),
            }

        # --- listing → four canonical buckets ---
        self._listing: dict[str, list[dict]] = self._build_listing(
            sheets[schema.LISTING_SHEET]
        )

    # -- construction helpers ------------------------------------------------

    def _build_documents(self, rows: list[dict]) -> None:
        """Group line rows into canonical documents, preserving first-seen order."""
        # key (doc_type, doc_kind, DocNum) → {"header": {...}, "lines": [(idx, line)]}
        grouped: dict[tuple, dict] = {}
        for row in rows:
            doc_type = schema.to_str(row.get(schema.DOC_TYPE_COL))
            doc_kind = schema.to_str(row.get(schema.DOC_KIND_COL))
            doc_num = schema.to_int(row.get("DocNum"))
            key = (doc_type, doc_kind, doc_num)
            entry = grouped.get(key)
            if entry is None:
                entry = {
                    "header": {
                        "DocNum": doc_num,
                        "DocDate": schema.to_str(row.get("DocDate")),
                        "CardCode": schema.to_str(row.get("CardCode")),
                        "CardName": schema.to_str(row.get("CardName")),
                        "DocCurrency": schema.to_str(row.get("DocCurrency")),
                        # Mirrors project_document — the round-trip asserts the two
                        # produce identical document shapes, so this must match.
                        "DocTotal": schema.to_float(row.get("DocTotal")),
                    },
                    "lines": [],
                }
                grouped[key] = entry
            line_index = schema.to_int(row.get(schema.LINE_INDEX_COL))
            entry["lines"].append((
                0 if line_index is None else line_index,
                {
                    "VatGroup": schema.to_str(row.get("VatGroup")),
                    "LineTotal": schema.to_float(row.get("LineTotal")),
                    "TaxTotal": schema.to_float(row.get("TaxTotal")),
                },
            ))

        for (doc_type, doc_kind, _doc_num), entry in grouped.items():
            lines = [ln for _idx, ln in sorted(entry["lines"], key=lambda t: t[0])]
            doc = dict(entry["header"])
            doc["DocumentLines"] = lines
            if doc_kind == "credit_note":
                doc["is_credit_note"] = True
            self._docs_by_bucket[(doc_type, doc_kind)].append(doc)

    def _build_listing(self, rows: list[dict]) -> dict[str, list[dict]]:
        buckets: dict[str, list[dict]] = {name: [] for name in schema.LISTING_BUCKETS}
        for row in rows:
            scope = schema.to_str(row.get(schema.SCOPE_COL))
            doc_type = schema.to_str(row.get(schema.DOC_TYPE_COL))
            for name, (want_scope, want_type, projector) in schema.LISTING_BUCKETS.items():
                if scope == want_scope and doc_type == want_type:
                    buckets[name].append(projector(row))
                    break
        return buckets

    # -- ChainReader surfaces (structural) -----------------------------------

    def count(self, entity: str, period_start: str, period_end: str) -> Optional[int]:
        """S0 — TRUE document count for the entity (the export has every row).

        Deliberately NOT the dormant ``@odata.count`` → None the live feeder/oracle
        reproduce: an export can count honestly. Never assert this against the frozen S0.
        """
        route = schema.ENTITY_ROUTING.get(entity)
        if route is None:
            return None
        return len(self._docs_by_bucket[route])

    def fetch_invoices(self, entity: str, period_start: str, period_end: str) -> list:
        """S1 — full line-level documents for ``entity`` (Invoices/PurchaseInvoices)."""
        doc_type, doc_kind = schema.ENTITY_ROUTING[entity]
        return copy.deepcopy(self._docs_by_bucket[(doc_type, doc_kind)])

    def fetch_credit_notes(self, entity_type: str, period_start: str, period_end: str) -> list:
        """S2 — credit notes for 'sales'/'purchases', tagged is_credit_note=True."""
        doc_type = schema.CREDIT_NOTE_TYPE_ROUTING[entity_type]
        return copy.deepcopy(self._docs_by_bucket[(doc_type, "credit_note")])

    def get_business_partner(self, card_code: str) -> dict:
        """S3 — single BusinessPartner master record by CardCode (carries FederalTaxID)."""
        if card_code not in self._bp_by_card:
            raise KeyError(f"CardCode {card_code!r} not present in export business partners")
        return copy.deepcopy(self._bp_by_card[card_code])

    def fetch_listing(self, period: dict) -> dict:
        """S5 — the four header-only listing buckets."""
        return copy.deepcopy(self._listing)

    # -- coverage seam (beside the Protocol; NOT part of ChainReader) ---------

    def coverage(self) -> ExtractCoverage:
        """Declare which canonical fields the export carried, header AND value (emission)."""
        fields = {
            key: (column in self._present_columns.get(sheet, set()))
            for key, (sheet, column) in schema.COVERAGE_FIELDS.items()
        }
        populated = {
            key: (column in self._populated_columns.get(sheet, set()))
            for key, (sheet, column) in schema.COVERAGE_FIELDS.items()
        }
        return ExtractCoverage(fields=fields, populated=populated)

    def populatable_sides(self) -> frozenset:
        """Which F5 SIDES this FORMAT can populate — a capability fact, never a content fact.

        The three-sheet extract format carries a ``doc_type`` column spanning sales AND
        purchase documents, so both sides are declarable even when a given export happens
        to contain only one — capability, not emptiness, drives the render marker (open
        item #48). Declared beside ``coverage_status()`` (same duck-typed seam family).
        """
        return frozenset({"sales", "purchase"})

    def coverage_status(self) -> list:
        """Map this export's coverage onto per-check status (T2.12 slice 2B + ext-1 + ext-3).

        Returns ``list[CoverageStatus]`` for the in-scope checks: the three locked 2B
        cases (DUP_CLAIM, NO_GST_REG, SEQ_GAP), then the four line-level E-checks (E1–E4,
        ext-1), then the four document-pre-pass checks (ext-3). SEQ_GAP's company-wide
        signal is the presence of any company-wide ('all' scope) sales rows — the surface
        ``detect_seq_gaps`` consumes to tell "issued in another period" from "never issued
        anywhere". An extract is a listing/transaction export with NO source-document (PDF)
        surface — the PDF provider is a separate engine-level seam this reader does not carry
        — so ``document_pdfs_present=False`` and the four document checks are ``unavailable``.
        Emission only; asserts no verdict.
        """
        company_wide_present = bool(self._listing.get("all_sales_headers"))
        return derive_coverage_statuses(
            self.coverage(),
            company_wide_population_present=company_wide_present,
            document_pdfs_present=False,
        )


# ---------------------------------------------------------------------------
# Loading — CSV directory or .xlsx workbook. Returns (sheets, present_columns)
# where sheets maps base-name → list[row-dict] and present_columns maps
# base-name → set of header columns actually found (for the coverage seam).
# ---------------------------------------------------------------------------

_REQUIRED_SHEETS = (
    schema.DOCUMENTS_SHEET,
    schema.BUSINESS_PARTNERS_SHEET,
    schema.LISTING_SHEET,
)


def _compute_populated_columns(sheets: dict) -> dict:
    """Per sheet, the set of columns that carried at least one NON-EMPTY value.

    A cell counts as populated when it stringifies to a non-blank value; whitespace-only
    cells (and ``None``) do not. This is what makes coverage value-population-aware: a column
    present in the header but empty in every row is present but NOT populated.
    """
    populated: dict[str, set] = {}
    for name, rows in sheets.items():
        cols: set = set()
        for row in rows:
            for col, val in row.items():
                if col in cols:
                    continue
                if val is not None and str(val).strip() != "":
                    cols.add(col)
        populated[name] = cols
    return populated


def _load_sheets(source: Path):
    if source.is_dir():
        return _load_csv_dir(source)
    if source.suffix.lower() == ".xlsx":
        return _load_xlsx(source)
    raise ValueError(
        f"unsupported export source {source!r}: expected a CSV directory or a .xlsx file"
    )


def _load_csv_dir(source: Path):
    sheets: dict[str, list[dict]] = {}
    present_columns: dict[str, set] = {}
    for name in _REQUIRED_SHEETS:
        path = source / f"{name}.csv"
        if not path.exists():
            raise FileNotFoundError(f"export missing {name}.csv at {path}")
        with path.open("r", encoding="utf-8", newline="") as fh:
            reader = csv.DictReader(fh)
            present_columns[name] = set(reader.fieldnames or [])
            sheets[name] = [dict(row) for row in reader]
    return sheets, present_columns


def _load_xlsx(source: Path):
    # Lazy import: the CSV path needs no third-party dependency.
    from openpyxl import load_workbook

    wb = load_workbook(source, read_only=True, data_only=True)
    sheets: dict[str, list[dict]] = {}
    present_columns: dict[str, set] = {}
    try:
        for name in _REQUIRED_SHEETS:
            if name not in wb.sheetnames:
                raise FileNotFoundError(f"export workbook missing sheet {name!r}")
            ws = wb[name]
            rows_iter = ws.iter_rows(values_only=True)
            try:
                header = list(next(rows_iter))
            except StopIteration:
                header = []
            header = [("" if h is None else str(h)) for h in header]
            present_columns[name] = set(header)
            records = []
            for raw in rows_iter:
                if all(cell is None for cell in raw):
                    continue
                # openpyxl yields native cell types; ser() re-stringifies so the xlsx
                # and CSV paths parse through the identical coercers.
                record = {
                    col: schema.ser(raw[i]) if i < len(raw) else ""
                    for i, col in enumerate(header)
                }
                records.append(record)
            sheets[name] = records
    finally:
        wb.close()
    return sheets, present_columns
