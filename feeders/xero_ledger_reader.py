"""feeders/xero_ledger_reader.py — load_gst_ledger: side-input loader for the Xero GST
control-account (820) ledger export ("Account Transactions" -> "GST Transactions" sheet).

This is a SIDE-INPUT LOADER, like orchestrator.check_declared_f5.load_declared_f5 — it is
NOT a ChainReader and does NOT implement or widen the ChainReader Protocol. It returns a flat
list of line dicts that the T2.24 internal-consistency check
(orchestrator.check_gst_ledger_recon) recomputes GST totals from.

Real-format rule (mirrors feeders.xero_f5_reader._locate_header): the export carries a title
block, so the column header is located STRUCTURALLY — the row whose cells include the markers
{"Date","Debit","Credit"} — never by a hard-coded row offset.

The non-transaction rows are skipped: the section label ("GST"), "Opening Balance",
"Total GST", "Closing Balance", "Total", and any fully-blank spacer. The "Total GST" /
"Closing Balance" rows are a DISPLAY ARTIFACT that does NOT foot cleanly under accrual +
unpaid invoices (in the real fixture "Total GST" reads 0); this loader NEVER reads them for
any total — every total is recomputed downstream from line-level Debit/Credit.

Pure stdlib; openpyxl is imported LAZILY inside the load function only. No SAP, no anthropic,
imports nothing upward (feeders is a leaf).
"""
from __future__ import annotations

from pathlib import Path
from typing import Any, Optional

# The sheet the Xero Account Transactions export carries the 820 control-account lines on.
GST_TRANSACTIONS_SHEET = "GST Transactions"

# Header column labels; the header row is the one carrying all three markers.
_COL_DATE = "Date"
_COL_SOURCE = "Source"
_COL_DESCRIPTION = "Description"
_COL_REFERENCE = "Reference"
_COL_DEBIT = "Debit"
_COL_CREDIT = "Credit"
_HEADER_MARKERS = (_COL_DATE, _COL_DEBIT, _COL_CREDIT)

# Non-transaction rows (identified by their leading label cell) that carry no ledger line.
_SKIP_LABELS = frozenset({
    "GST", "Opening Balance", "Total GST", "Closing Balance", "Total",
})


def _to_float(raw: Any) -> float:
    """Coerce a Debit/Credit cell to float; blank / None -> 0.0."""
    if raw is None or raw == "":
        return 0.0
    return float(raw)


def _locate_header(rows: list[list]) -> dict:
    """Find the header row structurally (carries all _HEADER_MARKERS) -> {label: col-index}.

    Includes ``_header_row_index`` so the caller iterates only the rows below it. Raises
    ValueError if no such row exists (an honest client-input failure, never a silent []).
    """
    for idx, raw in enumerate(rows):
        labels = ["" if c is None else str(c).strip() for c in raw]
        if all(marker in labels for marker in _HEADER_MARKERS):
            col = {label: pos for pos, label in enumerate(labels) if label}
            col["_header_row_index"] = idx
            return col
    raise ValueError(
        f"GST ledger header row not found (no row carrying all of {_HEADER_MARKERS!r})"
    )


def _cell(cells: list, index: Optional[int]) -> Any:
    """Safe positional cell access (rows may be short); missing -> None."""
    if index is None or index >= len(cells):
        return None
    return cells[index]


def load_gst_ledger(path) -> list[dict]:
    """Parse the 'GST Transactions' sheet into a flat list of ledger-line dicts.

    Each transaction line -> {date, source, description, reference, debit: float, credit: float}.
    ``date`` is kept as-is (may be a datetime). ``debit``/``credit`` are floats (blank -> 0.0).

    Skips the section label / balance / total rows (see module docstring); those are never
    read for any total. Raises ValueError if the sheet or its header row is absent.
    """
    source = Path(path)
    if source.suffix.lower() != ".xlsx":
        raise ValueError(
            f"unsupported GST ledger source {source!r}: expected a .xlsx workbook"
        )
    from openpyxl import load_workbook  # lazy — keep the module dependency-light

    wb = load_workbook(source, read_only=True, data_only=True)
    try:
        if GST_TRANSACTIONS_SHEET not in wb.sheetnames:
            raise ValueError(
                f"GST ledger workbook missing sheet {GST_TRANSACTIONS_SHEET!r}"
            )
        ws = wb[GST_TRANSACTIONS_SHEET]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()

    col = _locate_header(rows)
    lines: list[dict] = []
    for raw in rows[col["_header_row_index"] + 1:]:
        cells = list(raw)
        if not any((c is not None and str(c).strip() != "") for c in cells):
            continue  # blank spacer row
        label = str(_cell(cells, col.get(_COL_DATE)) or "").strip()
        if label in _SKIP_LABELS:
            continue  # section label / opening / total / closing row — never a ledger line
        lines.append({
            "date": _cell(cells, col.get(_COL_DATE)),
            "source": (str(_cell(cells, col.get(_COL_SOURCE)) or "").strip()),
            "description": (str(_cell(cells, col.get(_COL_DESCRIPTION)) or "").strip()),
            "reference": (str(_cell(cells, col.get(_COL_REFERENCE)) or "").strip()),
            "debit": _to_float(_cell(cells, col.get(_COL_DEBIT))),
            "credit": _to_float(_cell(cells, col.get(_COL_CREDIT))),
        })
    return lines
